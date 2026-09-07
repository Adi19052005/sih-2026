import csv
from pathlib import Path

from docx import Document
from pptx import Presentation


def generate_word_doc(
    title: str,
    content: str,
    filepath: str
) -> str:
    """
    Generate a Microsoft Word (.docx) document locally.
    """

    output_path = Path(filepath).resolve()

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document()

    document.add_heading(title, level=0)

    # Preserve paragraphs from generated content
    paragraphs = content.split("\n")

    for paragraph in paragraphs:
        paragraph = paragraph.strip()

        if paragraph:
            document.add_paragraph(paragraph)

    document.save(output_path)

    return str(output_path)


def generate_powerpoint(
    title: str,
    bullet_points: list,
    filepath: str
) -> str:
    """
    Generate a Microsoft PowerPoint (.pptx) presentation locally.
    """

    output_path = Path(filepath).resolve()

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    presentation = Presentation()

    # -------------------------
    # Slide 1: Title Slide
    # -------------------------

    title_slide_layout = presentation.slide_layouts[0]

    slide = presentation.slides.add_slide(
        title_slide_layout
    )

    slide.shapes.title.text = title

    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = (
            "Generated locally by Sovereign AI Workbench"
        )

    # -------------------------
    # Slide 2: Content Slide
    # -------------------------

    content_slide_layout = presentation.slide_layouts[1]

    slide = presentation.slides.add_slide(
        content_slide_layout
    )

    slide.shapes.title.text = "Key Points"

    text_frame = slide.placeholders[1].text_frame

    text_frame.clear()

    if bullet_points:

        for index, point in enumerate(bullet_points):

            if index == 0:
                paragraph = text_frame.paragraphs[0]
            else:
                paragraph = text_frame.add_paragraph()

            paragraph.text = str(point)
            paragraph.level = 0

    else:
        text_frame.paragraphs[0].text = (
            "No bullet points provided."
        )

    presentation.save(output_path)

    return str(output_path)


def generate_csv(
    headers: list,
    rows: list,
    filepath: str
) -> str:
    """
    Generate a CSV file locally.
    """

    output_path = Path(filepath).resolve()

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with open(
        output_path,
        mode="w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.writer(file)

        writer.writerow(headers)

        for row in rows:
            writer.writerow(row)

    return str(output_path)